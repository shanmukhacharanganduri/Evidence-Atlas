"""DOCX and PDF evidence packs built from an already-computed coverage result.

Exports consume the same rows the screen shows (no re-running of retrieval) and inherit the
selected institution scope and each source's synthetic status.
"""
from __future__ import annotations

import io
import textwrap
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape

FRAMEWORK = "NAAC / NBA"
_FONT_CANDIDATES = [
    Path("C:/Windows/Fonts/arial.ttf"),
    Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    Path("/Library/Fonts/Arial Unicode.ttf"),
]


def _header_line(scope: str, generated: datetime | None) -> str:
    stamp = (generated or datetime.now()).strftime("%d %B %Y, %H:%M")
    return f"Generated: {stamp}  |  Scope: {scope}  |  Framework: {FRAMEWORK}"


def _rows(coverage: list[dict], criterion: str | None) -> list[dict]:
    return [row for row in coverage if criterion is None or row["Criterion"] == criterion]


def _status_line(row: dict) -> str:
    note = " (synthetic demonstration source)" if row.get("synthetic") else ""
    return f"Status: {row['Evidence strength']} · {row.get('Review', 'Not reviewed')}{note}"


def export_docx(coverage: list[dict], scope: str, criterion: str | None = None,
                generated: datetime | None = None) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import RGBColor

    doc = Document()
    doc.add_heading("Accreditation Evidence Pack", 0).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(_header_line(scope, generated)).alignment = WD_ALIGN_PARAGRAPH.CENTER
    colours = {"strong": RGBColor(0x1B, 0x6E, 0x3E), "thin": RGBColor(0x92, 0x59, 0x00), "none": RGBColor(0x9B, 0x1C, 0x1C)}
    for row in _rows(coverage, criterion):
        heading = doc.add_heading(row["Criterion"], level=1)
        if heading.runs:
            heading.runs[0].font.color.rgb = colours.get(row.get("Color", ""), RGBColor(0, 0, 0))
        doc.add_paragraph().add_run(_status_line(row)).bold = True
        if row.get("Finding"):
            doc.add_paragraph(f"Finding: {row['Finding']}")
        if row.get("Citation") and row["Citation"] != "—":
            doc.add_paragraph().add_run(f"Citation: {row['Citation']}").italic = True
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _pdf_font() -> str:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for path in _FONT_CANDIDATES:
        if path.exists():
            try:
                pdfmetrics.registerFont(TTFont("EvidenceFont", str(path)))
                return "EvidenceFont"
            except Exception:
                continue
    return "Helvetica"


def export_pdf(coverage: list[dict], scope: str, criterion: str | None = None,
               generated: datetime | None = None) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer

    font = _pdf_font()
    styles = getSampleStyleSheet()
    normal = ParagraphStyle("N", parent=styles["Normal"], fontName=font)
    title = ParagraphStyle("T", parent=styles["Title"], fontName=font, fontSize=20, spaceAfter=4)
    colour = {"strong": "#1B6E3E", "thin": "#925900", "none": "#9B1C1C"}
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm, bottomMargin=2 * cm)
    story = [Paragraph("Accreditation Evidence Pack", title),
             Paragraph(escape(_header_line(scope, generated)), normal), Spacer(1, 0.5 * cm)]
    for row in _rows(coverage, criterion):
        heading = ParagraphStyle("H", parent=styles["Heading2"], fontName=font, spaceBefore=10,
                                 textColor=colors.HexColor(colour.get(row.get("Color", ""), "#000000")))
        story.append(Paragraph(escape(row["Criterion"]), heading))
        story.append(Paragraph(f"<b>{escape(_status_line(row))}</b>", normal))
        if row.get("Finding"):
            story.append(Paragraph(f"Finding: {escape(textwrap.shorten(row['Finding'], width=400, placeholder='…'))}", normal))
        if row.get("Citation") and row["Citation"] != "—":
            story.append(Paragraph(f"<i>Citation: {escape(row['Citation'])}</i>", normal))
        story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#cac7be"), spaceAfter=6))
    doc.build(story)
    return buffer.getvalue()
